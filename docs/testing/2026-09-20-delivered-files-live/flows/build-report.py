from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(sys.argv[1])


def set_cell_fill(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def set_cell_borders(cell, color="D9D9D9"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)
        borders.append(element)


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = document.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(7)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(25)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
title_properties = styles["Title"].element.get_or_add_pPr()
title_borders = title_properties.find(qn("w:pBdr"))
if title_borders is not None:
    title_properties.remove(title_borders)
for name, size in (("Heading 1", 16), ("Heading 2", 12)):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

title = document.add_paragraph(style="Title")
title.add_run("Delivered Files Acceptance Report")
subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = subtitle.add_run("VUH 1105 implementation and verification scope")
run.bold = True
run.font.size = Pt(11)

document.add_paragraph(
    "Clankie now publishes deliberate conversation deliverables through a bounded file store and returns exact bytes only to authenticated callers. The path uses the existing conversation log, attachment storage root, device grants, relay, and native file handling. This report records the implemented contract and the proof required before the issue can close."
)

document.add_heading("Delivery contract", level=1)
contract = document.add_table(rows=1, cols=3)
contract.autofit = False
widths = (Inches(1.55), Inches(2.25), Inches(3.85))
headers = ("Boundary", "Contract", "Result")
for index, header in enumerate(headers):
    cell = contract.rows[0].cells[index]
    cell.width = widths[index]
    cell.text = header
    set_cell_fill(cell, "18324A")
    for paragraph in cell.paragraphs:
        for item in paragraph.runs:
            item.font.bold = True
            item.font.color.rgb = RGBColor(255, 255, 255)
    set_cell_borders(cell)
rows = [
    ("Publication", "CLI and captain tool", "An existing authorized regular file becomes a durable file event with name, content type, byte size, hash, and stable artifact reference."),
    ("Storage", "Private content addressed file", "Bytes live under the existing attachment root. Atomic writes use a unique pending filename so concurrent publication of identical content does not collide."),
    ("Retrieval", "POST /operator/v1/artifacts/download", "The request body carries conversationId and artifactId. The response carries raw bytes and stored response headers. Credentials never enter the URL."),
    ("Authorization", "Operator or device chat grant", "The captain route requires operator authentication. The relay validates the current device grant before requesting bytes upstream."),
    ("Consumer", "Encrypted fetch then native local file", "The app downloads through its injected transport, writes protected local bytes, and opens the system preview and share surfaces."),
]
for row_values in rows:
    cells = contract.add_row().cells
    for index, value in enumerate(row_values):
        cells[index].width = widths[index]
        cells[index].text = value
        cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_borders(cells[index])

document.add_heading("Trust boundaries", level=1)
for text in [
    "Publication resolves the source through realpath and accepts only regular files inside the conversation workspace.",
    "A file may not exceed 15 MiB. Names and content types pass explicit validation before storage.",
    "The artifact download uses POST on a strict relative host route. No bearer, ticket, capability, or download secret appears in a URL.",
    "The relay refuses remote path publication and checks the current chat grant for every download.",
    "Stored metadata and bytes remain hash bound. Missing, mismatched, revoked, unrelated, or unauthenticated requests fail closed.",
]:
    document.add_paragraph(text, style="List Bullet")

document.add_heading("Four sourced deliverables", level=1)
deliverables = document.add_table(rows=1, cols=4)
deliverables.autofit = False
deliverable_widths = (Inches(1.35), Inches(1.35), Inches(2.35), Inches(2.6))
for index, header in enumerate(("Type", "Format", "Source", "Use")):
    cell = deliverables.rows[0].cells[index]
    cell.width = deliverable_widths[index]
    cell.text = header
    set_cell_fill(cell, "DCE6F1")
    for paragraph in cell.paragraphs:
        for item in paragraph.runs:
            item.font.bold = True
    set_cell_borders(cell)
deliverable_rows = [
    ("Report", "DOCX", "ADR 0174, CLI docs, route and store code", "Readable implementation and acceptance record"),
    ("Spreadsheet", "XLSX", "Protocol constants and consumer behavior", "Scannable artifact and boundary matrix"),
    ("Presentation", "PPTX", "The same delivery contract and proof plan", "Short stakeholder walkthrough"),
    ("Website bundle", "ZIP", "Static HTML, CSS, JSON, README and source list", "Portable site that can be unpacked and opened locally"),
]
for row_values in deliverable_rows:
    cells = deliverables.add_row().cells
    for index, value in enumerate(row_values):
        cells[index].width = deliverable_widths[index]
        cells[index].text = value
        cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_borders(cells[index])

document.add_heading("Verification plan", level=1)
document.add_paragraph(
    "The isolated run publishes all four files into one fresh conversation, downloads each through the authenticated relay, compares byte count and SHA 256, restarts the host with the same state, and repeats retrieval. It also records refusal for missing authorization, an unrelated conversation, and an invalid source path. The simulator then pairs with that host and opens and shares a downloaded file through the production consumer path."
)

document.add_heading("Sources", level=1)
for source in [
    "docs/adr/0174-conversation-delivered-files.md",
    "docs/cli.md",
    "apps/clankie/src/delivered-files.ts",
    "apps/clankie/src/app.ts",
    "apps/relay/src/operator-conversations.ts",
    "packages/protocol/src/public-gateway.ts",
    "clankie-app packages/command-center/src/composer/liveCaptainTransport.ts",
    "clankie-app apps/mobile/modules/expo-clankie-files",
    "Linear VUH-1105",
]:
    document.add_paragraph(source, style="List Bullet")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
